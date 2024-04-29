import streamlit as st
import pandas as pd
import re
from datetime import datetime
from docx import Document
from io import StringIO
from streamlit.logger import get_logger

LOGGER = get_logger(__name__)

# Function to extract attendance data
# @st.cache_data
def extract_attendance(file_obj, start_date, end_date, whatsapp_name):
    attendance_data = []
    time_in = None

    # Define patterns for clocking in and out
    patterns_in = [
        r'Morning, clocking in, hi {}',
        r'Morning...clocking in',
        r'Good morning, clocking in',
        r'morning clock in',
        r'Morning clock in',
        r'Morning... clocking in',
        r'Morning clocking in',
        r'Morning clocking in',
        r'Morning.. clocking in',
        r'Good morning, clock in',
        r'Aisyah, winnie, stefano clocked in \(away day\)',
        r'Mornin clocking in',
        r'morning clocking in',
        r'Mornin , clocking in',
        r'Morning, clocking in',
        r'morning, clocking in',
        r'Morning...Clock in',
        r'Morning...Clocking in \(WFH\)',
        r'Clocking in for wfh 👨\u200d💻',
        r'morning clocking in for wfh',
        r'Morning clock in for wfh',
        r'Clock in also wfh 😊',
        r'clock in \(wfh\)',
        r'Clocking in \(wfh\)',
        r'Clocking in \(wfh\)',
        r'Clocking in for wfh too',
        r'clock in'
    ]

    patterns_out = [
        r'Clocking out',
        r'Clocked out',
        r'Bye..clocking out',
        r'Clocking out 👍',
        r'Clocking out 🫡',
        r'Bye...clocking out',
        r'Clocking out bye',
        r'Bye... clocking out',
        r'Clocking out \(wfh\)',
        r'Clocking out...bye',
        r'Bye clocking out',
        r'Clocked out at 6.00PM 🦦',
        r'Clocking out byeee',
        r'Clocking out happy weekenddd✨',
        r'Clocking out for today🙏',
        r'alright, thank you miss. Clocking out \(WFH\)',
        r'I’m clocking out too',
        r'Clocked our',
        r'Clocking out \(wfh\)',
        r'Clocking out🙏',
        r'clock out \(wfh\)',
        r'clock out',
        r'Clocking  out'
    ]

    pattern_in = '|'.join(patterns_in)
    pattern_out = '|'.join(patterns_out)

    # Constructing the regex pattern with the whatsapp_name
    pattern_in = r'\[(.*?)\] {}:\s*(.*?\b(?:{})\b.*?)\s*'.format(whatsapp_name.lower(), pattern_in)
    pattern_out = r'\[(.*?)\] {}:\s*(.*?\b(?:{})\b.*?)\s*'.format(whatsapp_name.lower(), pattern_out)

    # Iterate through each line in the file
    for line in file_obj:
        # Check for clocking in
        match_in = re.search(pattern_in, line.lower(), re.IGNORECASE)
        if match_in:
            date_str, _ = match_in.groups()
            date = datetime.strptime(date_str, '%d/%m/%Y, %H:%M:%S')
            if start_date <= date.date() <= end_date:
                time_in = date.time()

        # Check for clocking out
        match_out = re.search(pattern_out, line.lower(), re.IGNORECASE)
        if match_out and time_in is not None:
            date_str, _ = match_out.groups()
            date = datetime.strptime(date_str, '%d/%m/%Y, %H:%M:%S')
            if start_date <= date.date() <= end_date:
                attendance_data.append({
                    'Date': date.date(),
                    'Time_In': time_in.strftime('%H:%M'),
                    'Time_Out': date.time().strftime('%H:%M')
                })
                time_in = None

    # Create a DataFrame from the attendance data
    attendance_df = pd.DataFrame(attendance_data)

    # Convert the 'Date' column to datetime64[ns] data type
    attendance_df['Date'] = pd.to_datetime(attendance_df['Date'])

    # Create a range of dates between start_date and end_date
    date_range = pd.date_range(start_date, end_date, freq='D')

    # Create a DataFrame with the date range
    date_range_df = pd.DataFrame({'Date': date_range})

    # Merge the extracted attendance data with the date range DataFrame
    merged_df = date_range_df.merge(attendance_df, how='left', on='Date')

    # Remove the time component from the Date column
    merged_df['Date'] = merged_df['Date'].dt.date

    # Fill missing values with 'Holiday'
    merged_df = merged_df.fillna({'Time_In': 'Holiday', 'Time_Out': 'Holiday'})

    return merged_df

# Streamlit app
def app():
    st.set_page_config(
        page_title="Attendance Sheet Automation System (ASAS)",
        page_icon="👍",
    )

    st.title("Attendance Sheet Automation System (ASAS)")

    # File upload
    file = st.file_uploader("Upload WhatsApp chat file", type="txt")

    # Input fields
    start_date = st.date_input("Start Date", value=datetime(2024, 3, 1).date())
    end_date = st.date_input("End Date", value=datetime(2024, 3, 31).date())
    whatsapp_name = st.text_input("WhatsApp Name", value="farysa")

    if file:
        # Read the contents of the uploaded file
        file_content = file.getvalue().decode("utf-8")

        # Create a StringIO object with the file content
        file_stream = StringIO(file_content)

        # Extract attendance data
        attendance_df = extract_attendance(file_stream, start_date, end_date, whatsapp_name)

        # Display attendance data
        st.write("Attendance Data:")
        st.dataframe(attendance_df)

        st.write("Trainee Details:")
        # Get user input
        name = st.text_input("Name", value="Amirul Asyraf")
        dept_div = st.text_input("Department/Division", value="GTS-Energy")
        phone_number = st.text_input("Phone Number", value="O13-35097092")
        email = st.text_input("Email", value="amirulasyraf.zainala@petronas.com")
        ic_number = st.text_input("IC Number", value="011102-14-0551")
        cost_centre = st.text_input("Cost - Centre", value="001PCT-606")
        ext_no = st.text_input("EXT NO")
        contact_person = st.text_input("Contact Person", value="Rafidah Bt. Ramli")

        # start_date = st.text_input("Start Date")
        # end_date = st.text_input("End Date")
        # whatsapp_name = st.text_input("Whatsapp name")

        position = st.text_input("Position", value="Intern")
        # date_today = st.date_input("Date")
        date_today = st.text_input("Date Today", value="2024/03/01")
        sv_name = st.text_input("Name (Supervisor)", value="Tanasegaran Kuppusamy")
        sv_position = st.text_input("Position (Supervisor)", value="Staff (Energy)")
        # sv_date = st.text_input("Date (Supervisor)")

        # Export to Word document
        if st.button("Export to Word"):
            doc = Document('Attendance_Sheet.docx')
            table = doc.tables[0]

            # Fill in personal details
            table.cell(2, 2).text = name
            table.cell(2, 9).text = ic_number
            table.cell(3, 2).text = dept_div
            table.cell(3, 9).text = cost_centre
            table.cell(4, 2).text = phone_number
            table.cell(4, 9).text = ext_no
            table.cell(5, 2).text = email
            table.cell(5, 10).text = contact_person
            table.cell(44, 2).text = name
            table.cell(44, 9).text = sv_name
            table.cell(45, 2).text = position
            table.cell(45, 9).text = sv_position
            table.cell(46, 2).text = date_today

            # Write attendance data into the table
            for i, row in attendance_df.iterrows():
                # print(f"i : {i} , row : {row}")
                for j, value in enumerate(row):
                    # print(f"j : {j} , value : {value}")
                    table.cell(i + 10, j + 3).text = str(value)

            # Save the modified document
            modified_doc_path = 'modified_document.docx'
            doc.save(modified_doc_path)
            st.success("Attendance data exported to Word document successfully!")

            # Provide download link for the modified document
            with open(modified_doc_path, "rb") as file:
                btn = st.download_button(
                    label="Download Modified Document",
                    data=file,
                    file_name="modified_document.docx",
                    mime="application/octet-stream"
                )
    else:
        st.warning("Please upload a WhatsApp chat file.")

if __name__ == "__main__":
    app()